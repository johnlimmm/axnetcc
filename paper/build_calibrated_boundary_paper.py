from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "MNC_FLOW_경계제약형_전문가라우팅_국내학술대회_논문.docx"
FIG_ARCH = ROOT / "architecture.png"
FIG_PARETO = ROOT / "calibrated_pareto.png"
FONT = "Malgun Gothic"
INK = "1F2937"
BLUE = "244A78"
LIGHT = "EAF0F7"
GRID = "AAB4C2"


def set_font(run, size=8.4, bold=False, color=INK, italic=False):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), FONT)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def para(doc, text, size=8.4, after=3, first=3.5, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.13
    p.paragraph_format.first_line_indent = Mm(first)
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), size=size, bold=True, color=BLUE)
        set_font(p.add_run(text[len(bold_prefix):]), size=size)
    else:
        set_font(p.add_run(text), size=size)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    set_font(
        p.add_run(text),
        size=10.5 if level == 1 else 9.2,
        bold=True,
        color=BLUE if level == 1 else "1F3A5F",
    )
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), size=7.2, color="4B5563")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for name, value in (("top", 70), ("bottom", 70), ("start", 100), ("end", 100)):
                margin = tc_mar.find(qn(f"w:{name}"))
                if margin is None:
                    margin = OxmlElement(f"w:{name}")
                    tc_mar.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths, highlight_last=True):
    tbl = doc.add_table(rows=1, cols=len(headers))
    for index, text in enumerate(headers):
        shade(tbl.rows[0].cells[index], BLUE)
        p = tbl.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=6.8, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = tbl.add_row().cells
        for index, value in enumerate(values):
            if highlight_last and row_index == len(rows) - 1:
                shade(cells[index], LIGHT)
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(str(value)), size=6.6, bold=highlight_last and row_index == len(rows) - 1)
    set_table_geometry(tbl, widths)


def page_break(doc):
    doc.add_page_break()


def create_pareto():
    points = {
        "Static": (2092, 72.7),
        "Top-k": (2138, 71.9),
        "MasRouter-adapted": (1879, 68.9),
        "RouteLLM-adapted": (1809, 67.4),
        "IRT-adapted": (2023, 67.6),
        "Learned": (1437, 66.4),
        "Calibrated Boundary": (1607, 69.8),
    }
    width, height = 1500, 560
    left, right, top, bottom = 120, 40, 40, 90
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    font = ImageFont.truetype(str(font_path), 25) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 20) if font_path.exists() else ImageFont.load_default()
    xmin, xmax, ymin, ymax = 1350, 2200, 65, 74
    sx = lambda value: left + (value - xmin) / (xmax - xmin) * (width - left - right)
    sy = lambda value: height - bottom - (value - ymin) / (ymax - ymin) * (height - top - bottom)
    for tick in [1400, 1600, 1800, 2000, 2200]:
        x = sx(tick)
        draw.line((x, top, x, height - bottom), fill="#D7DEE7", width=1)
        draw.text((x - 25, height - bottom + 12), str(tick), fill=f"#{INK}", font=small)
    for tick in [66, 68, 70, 72, 74]:
        y = sy(tick)
        draw.line((left, y, width - right, y), fill="#D7DEE7", width=1)
        draw.text((35, y - 12), str(tick), fill=f"#{INK}", font=small)
    draw.line((left, top, left, height - bottom), fill=f"#{INK}", width=2)
    draw.line((left, height - bottom, width - right, height - bottom), fill=f"#{INK}", width=2)
    for name, (x_value, y_value) in points.items():
        x, y = sx(x_value), sy(y_value)
        accent = name == "Calibrated Boundary"
        radius = 10 if accent else 7
        color = "#C73E1D" if accent else "#365F8D"
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color)
        draw.text((x + 10, y - 25), name, fill=color if accent else f"#{INK}", font=small)
    draw.text((width // 2 - 210, height - 40), "Average cross-boundary payload proxy (bytes)", fill=f"#{INK}", font=font)
    draw.text((10, 8), "Micro-F1 (%)", fill=f"#{INK}", font=font)
    image.save(FIG_PARETO)


def create_architecture():
    width, height = 1600, 390
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/malgun.ttf")
    font = ImageFont.truetype(str(font_path), 27) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 22) if font_path.exists() else ImageFont.load_default()
    boxes = [
        (45, 115, 275, 275, "사용자 질의\n+ 기관 정책"),
        (355, 85, 670, 305, "경계 제약\n· 반출 금지 규칙\n· 필수 전문가 고정"),
        (750, 85, 1065, 305, "보정 라우터\n· 규칙/학습/유사도\n· 선택적 fan-out"),
        (1145, 55, 1555, 335, "독립 전문가 에이전트\n법률 · 보안 · 조달 · 행정\n로컬 LLM + 기관별 RAG"),
    ]
    for index, (x1, y1, x2, y2, label) in enumerate(boxes):
        fill = "#EAF0F7" if index in (1, 2) else "#F7F9FC"
        outline = "#244A78" if index in (1, 2) else "#6B7C93"
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=fill, outline=outline, width=4)
        lines = label.split("\n")
        heights = [draw.textbbox((0, 0), line, font=font if i == 0 else small)[3] for i, line in enumerate(lines)]
        y = (y1 + y2 - sum(heights) - 5 * (len(lines) - 1)) / 2
        for i, line in enumerate(lines):
            active_font = font if i == 0 else small
            bbox = draw.textbbox((0, 0), line, font=active_font)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, fill=f"#{INK}", font=active_font)
            y += heights[i] + 5
    for x1, x2 in ((275, 355), (670, 750), (1065, 1145)):
        y = 195
        draw.line((x1 + 8, y, x2 - 15, y), fill="#244A78", width=5)
        draw.polygon(((x2 - 15, y - 10), (x2 - 15, y + 10), (x2, y)), fill="#244A78")
    image.save(FIG_ARCH)


create_pareto()
create_architecture()
doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(13)
section.bottom_margin = Mm(13)
section.left_margin = Mm(16)
section.right_margin = Mm(16)
section.header_distance = Mm(6)
section.footer_distance = Mm(7)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(8.4)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.13
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("MNC-Flow | 국내 학술대회 투고 원고"), size=7, color="6B7280")

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
set_font(title.add_run("공공·기업 AX를 위한 경계 제약형 보정 전문가 라우팅"), size=16, bold=True, color="102A43")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
set_font(
    subtitle.add_run("Boundary-Constrained Calibrated Expert Routing for Public and Enterprise AX"),
    size=8.5,
    italic=True,
    color="4B5563",
)
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.paragraph_format.space_after = Pt(6)
set_font(author.add_run("임지훈 | 고려대학교 MNC Lab."), size=8.5, bold=True)

para(
    doc,
    "요약: 공공기관과 기업의 생성형 AI 질의는 기술·보안·법무·정책 등 복수 전문영역을 요구하지만, "
    "중앙집중형 RAG는 원문 이동과 권한 확대를 유발하고 All-Agent 방식은 불필요한 호출과 통신량을 증가시킨다. "
    "본 논문은 정확도 점수와 데이터 경계 정책을 분리하고, 경계 정책을 우회할 수 없는 hard constraint로 적용하는 "
    "경계 제약형 보정 전문가 라우터를 제안한다. 3개 조직 시나리오와 6개 업무 관점으로 구성된 240개 후보를 서로 다른 "
    "OpenAI mini 모델이 독립 라벨링하고, 불일치를 제3 모델이 공공문서 원문 근거로 조정하였다. 근거 조건을 통과한 "
    "226개 문항을 40개 의미군 단위로 분리한 5-fold 평가에서 제안 방식은 Micro-F1 69.8, 평균 fan-out 3.23, "
    "경계 전송량 proxy 1,607 B를 기록하였다. 최고 Static baseline 대비 품질 96.0%를 유지하면서 fan-out을 15.0%, "
    "전송량을 23.2% 줄였고, MasRouter·RouteLLM·IRT 구조 adapter보다 높은 F1과 낮은 전송량을 동시에 달성하였다. "
    "다만 5점 비열등성은 통계적으로 확증되지 않아 추가 독립 의미군과 사람 검증이 필요하다.",
    size=8.0,
    first=0,
    bold_prefix="요약:",
)
para(
    doc,
    "주요어: Multi-Agent RAG, Expert Routing, Data Boundary, Local LLM, Public AX, Pareto Optimization",
    size=8.0,
    first=0,
    bold_prefix="주요어:",
)

heading(doc, "1. 서론")
para(
    doc,
    "공공·기업 AX 질의는 하나의 전문영역에 머물지 않는다. 예를 들어 민원 AI 도입은 RAG 아키텍처뿐 아니라 "
    "개인정보 처리, 접근통제, 위탁계약, 설명가능성, 조달 및 운영 책임을 동시에 검토해야 한다. 중앙 Core가 모든 부서 "
    "문서를 수집하면 통합은 쉬워지지만 최소권한과 목적제한 원칙이 약화된다. 반대로 모든 전문 Agent를 호출하면 누락은 "
    "줄지만 불필요한 문서 접근, 모델 실행, 통신량이 증가한다."
)
para(
    doc,
    "기존 LLM routing은 주로 강·약 모델 선택과 비용 절감에 집중한다. RouteLLM은 선호 데이터로 모델 선택을 학습하고[4], "
    "IRT-Router는 모델 능력과 질의 난도를 결합한다[5]. MasRouter는 협업 방식, 역할, LLM 선택을 cascade로 구성하지만[3], "
    "조직별 문서 경계와 Core로 공개 가능한 출력 계약은 직접적인 최적화 대상이 아니다. RemoteRAG는 cloud query privacy를 "
    "다루지만[7] 복수 부서의 전문가 선택 문제와는 축이 다르다."
)
para(
    doc,
    "본 연구의 핵심 질문은 ‘최고 정확도의 router는 무엇인가’가 아니라 ‘허용 가능한 품질을 유지하면서 어떤 Agent가 어떤 "
    "경계 안에서 추론하고 무엇만 Core로 전달할 것인가’이다. 기여는 세 가지다. 첫째, 정확도 calibration과 데이터 경계 "
    "hard constraint를 결합한 전문가 선택기를 제시한다. 둘째, 공공문서 근거와 상용 LLM 보조 라벨링으로 226문항 평가 "
    "세트를 구축한다. 셋째, 품질·fan-out·경계 전송량을 함께 비교하고 Pareto 관점에서 실용성을 분석한다."
)

heading(doc, "2. 관련 연구와 연구 공백")
heading(doc, "2.1 RAG와 Multi-Agent 집계", 2)
para(
    doc,
    "RAG는 비모수 지식을 검색해 생성 모델의 근거성과 최신성을 보완한다[1]. Mixture-of-Agents는 여러 모델 출력을 계층적으로 "
    "집계해 품질 향상을 보였으나[2], 참여 Agent 수와 중간 출력 전달이 늘수록 실행·통신 비용이 커진다. 공공·기업 환경에서는 "
    "비용 외에도 권한 없는 문서 접근과 원문 이동이 독립적인 위험이다."
)
heading(doc, "2.2 학습형 Router와 보호형 RAG", 2)
para(
    doc,
    "RouterBench는 다중 LLM routing을 성능·비용 문제로 정식화하고 대규모 inference 결과를 제공한다[6]. RouteLLM과 "
    "IRT-Router는 모델 선택에서 강점을 보이며, MasRouter는 이를 Multi-Agent collaboration으로 확장한다. 그러나 모델 선택과 "
    "전문 역할 선택은 동일하지 않고, 원 논문 checkpoint를 소규모 한국어 AX 역할 선택에 직접 적용할 수 없다. 따라서 본 "
    "연구는 공식 구조를 보존한 AX adapter와 단순하지만 강한 Static·Top-k baseline을 함께 사용한다. NIST AI RMF가 강조하는 "
    "Govern·Map·Measure·Manage 관점[8]에서 routing 결정 자체도 감사 가능한 통제 지점이 되어야 한다."
)

heading(doc, "2.3 문제의 중요성과 구조적 난점", 2)
para(
    doc,
    "공공·기업 AX의 전문가 라우팅은 단순한 성능·비용 최적화 문제가 아니다. 하나의 업무 질의에는 개인정보, 내부 규정, 법률 해석, "
    "조달 절차와 기술 운영이 동시에 얽히며 각 정보는 서로 다른 조직·보안 경계에 놓인다. 중앙 Core가 모든 원문을 수집하면 답변 통합은 "
    "쉬워지지만 데이터 최소화와 목적 제한 원칙을 훼손할 수 있다. 반대로 경계를 엄격히 분리하면 필요한 전문가와 근거를 발견하지 못해 답변 "
    "품질과 책임성이 떨어진다. 결국 품질, 데이터 보호, 운영 효율을 동시에 만족해야 하는 구조적 긴장이 존재한다."
)
para(
    doc,
    "이 문제는 세 가지 이유로 어렵다. 첫째, 전문가 역할 선택은 하나의 질의에 복수 정답이 존재하는 multilabel 문제이다. 과소 선택은 "
    "필수 검토 누락을, 과다 선택은 불필요한 문서 접근과 전송을 만든다. 둘째, 개인정보나 법적 검토 같은 정책 의무는 학습 점수가 낮더라도 "
    "반드시 실행되어야 하므로 평균 정확도만 최적화하는 확률적 router에 맡길 수 없다. 셋째, 민감 원문을 중앙에서 모아 선택 정확도를 "
    "높이는 행위 자체가 보호하려는 데이터 경계를 침해한다. 즉, 라우터는 원문을 보지 않으면서도 충분한 전문가를 찾아야 한다."
)
para(
    doc,
    "따라서 본 연구의 목표는 최고 정확도 경쟁이 아니라 정책 위반 가능성을 구조적으로 차단한 상태에서 최고 품질 방식에 근접하고 호출 범위와 "
    "전송량을 줄이는 것이다. 보호 규칙을 사후 필터로 추가하는 대신 라우팅의 실행 가능 영역을 hard constraint로 제한하고, 남은 후보만 "
    "보정 점수로 선택한다. 이 설계는 운영기관이 품질과 위험 사이의 선택을 Pareto 관계로 설명하고 감사할 수 있게 한다."
)

heading(doc, "3. 경계 제약형 보정 전문가 라우팅")
doc.add_picture(str(FIG_ARCH), width=Mm(170))
caption(doc, "그림 1. Edge 전문 Agent와 Core 사이의 최소 공개형 오케스트레이션 구조")
para(
    doc,
    "각 Agent a_i는 허용 문서 집합 D_i, 전문 역할 R_i, 보안등급 S_i, 공개 가능한 출력 스키마 O_i를 갖는다. 원문과 검색 "
    "context는 Edge 경계를 벗어나지 않고, Core에는 decision, confidence, required_actions, evidence_handles, "
    "disclosure_level, unresolved_conflicts만 전달한다. 개인정보·기밀·권한 관련 질의는 정책상 필수 Agent를 강제한다."
)
para(
    doc,
    "질의 q와 Agent i에 대해 규칙 기반 선택 여부 p_i^rule, fold 내부에서 학습한 multilabel 확률 p_i^learn, 역할 profile "
    "유사도 s_i를 계산하고 다음 보정 점수를 사용한다.",
)
equation = doc.add_paragraph()
equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(equation.add_run("z_i(q) = w_r p_i^rule(q) + w_l p_i^learn(q) + (1-w_r-w_l)s_i(q)"), size=9, italic=True)
para(
    doc,
    "학습 fold에서 (w_r, w_l), threshold τ, 최대 fan-out m을 선택하고, 테스트에서는 고정한다. 먼저 z_i≥τ인 상위 m개를 "
    "선택한 후 개인정보·기밀 등급에 필요한 hard constraint Agent를 합집합한다. 본 구현은 calibration 목적함수에 작은 "
    "fan-out 벌점을 적용해 F1이 유사한 후보 중 더 희소한 경로를 선택한다. 이 구조는 경계 정책이 학습 점수에 의해 제거되지 "
    "않는다는 점에서 일반적인 cost-aware thresholding과 다르다."
)
heading(doc, "3.1 데이터셋과 LLM 보조 라벨링", 2)
para(
    doc,
    "NIA·KISA·개인정보보호위원회·조달 관련 자료에서 3,027개 RAG chunk를 구성하고 문서 ID, 출처 URL, 발행일, SHA-256을 "
    "기록하였다. 기존 40개 파일럿 문항을 기반으로 중앙행정기관·지자체·민간기업과 설계·감사·사고대응·조달·성과·확산 관점을 "
    "조합해 240개 후보를 만들었다. A는 gpt-5.4-mini, B는 gpt-5-mini로 독립 판정하고 역할 집합이 다르거나 직접 인용이 "
    "유효하지 않으면 gpt-5.4-mini가 조정하였다."
)
table(
    doc,
    ["항목", "결과", "해석"],
    [
        ["전체 후보", "240", "조직별 80, 업무 관점별 40"],
        ["근거 조건 통과", "226 (94.2%)", "14문항은 사람 검토 큐로 제외"],
        ["A/B Micro-F1", "79.5%", "부분 역할 일치는 높음"],
        ["A/B exact agreement", "18.8%", "상용 LLM을 사람 전문가로 간주 불가"],
        ["제3 모델 조정", "85.0%", "불일치 자체를 품질 정보로 보존"],
    ],
    [2200, 1600, 5500],
    highlight_last=False,
)
caption(doc, "표 1. OpenAI 보조 라벨링 결과")
para(
    doc,
    "평가 누출을 막기 위해 동일 파일럿에서 파생된 여섯 관점은 semantic_family_id로 묶어 같은 fold에 배치하였다. 따라서 "
    "관측치는 226문항이지만 통계적 독립 단위는 40개 의미군이다. LLM 조정 라벨은 개발·비교용이며 사람 전문가 gold label로 "
    "표현하지 않는다."
)
heading(doc, "3.2 비교 방식과 지표", 2)
para(
    doc,
    "Random, Static, semantic Top-k, threshold, Naive Bayes multilabel, cost-aware와 함께 MasRouter·RouteLLM-MF·"
    "IRT-Router 구조 adapter를 비교하였다. Adapter는 원 논문의 전체 checkpoint 재현이 아니라 역할 선택 문제로 대응한 "
    "구현이다. 지표는 Micro/Macro-F1, exact match, hamming loss, fan-out, 구조화 AgentOutput의 byte proxy이다. "
    "질의 반복이 아니라 의미군 단위 permutation과 bootstrap을 사용하고 다중 비교에는 Holm 보정을 적용하였다."
)

page_break(doc)
heading(doc, "4. 실험 결과")
table(
    doc,
    ["방식", "P", "R", "Micro-F1", "Fan-out", "Bytes"],
    [
        ["Static", "74.0", "71.5", "72.7", "3.80", "2,092"],
        ["Top-k", "71.2", "72.5", "71.9", "4.00", "2,138"],
        ["MasRouter-adapted", "68.1", "69.6", "68.9", "4.01", "1,879"],
        ["RouteLLM-adapted", "67.4", "67.5", "67.4", "3.93", "1,809"],
        ["IRT-Router-adapted", "63.8", "72.0", "67.6", "4.43", "2,023"],
        ["Learned multilabel", "76.8", "58.4", "66.4", "2.99", "1,437"],
        ["기존 Boundary", "68.6", "73.9", "71.1", "4.23", "2,541"],
        ["Calibrated Boundary", "77.3", "63.6", "69.8", "3.23", "1,607"],
    ],
    [2400, 900, 900, 1300, 1200, 1300],
)
caption(doc, "표 2. 226문항 family-disjoint 5-fold 전문가 선택 결과")
para(
    doc,
    "Static은 Micro-F1 72.7로 가장 높았고 Top-k가 71.9였다. Calibrated Boundary는 69.8로 최고값의 96.0%를 유지하면서 "
    "Static 대비 fan-out을 15.0%, byte proxy를 23.2% 줄였다. 기존 Boundary와 비교하면 Precision은 8.7%p 증가했고 "
    "fan-out과 전송량은 각각 23.6%, 36.8% 감소했다. 세 연구 기반 adapter와 비교해서도 F1은 0.9~2.4%p 높고 전송량은 "
    "11.2~20.6% 적었다."
)
doc.add_picture(str(FIG_PARETO), width=Mm(165))
caption(doc, "그림 2. 전문가 선택 품질과 경계 전송량의 Pareto 관계")
para(
    doc,
    "Calibrated Boundary는 다른 방식에 의해 품질과 전송량 양쪽에서 동시에 지배되지 않는 Pareto 전선에 위치했다. "
    "Threshold·cost-aware 방식은 더 적은 byte를 사용했지만 recall이 각각 30.4%, 40.3%에 머물렀다. 반대로 Static과 "
    "Top-k는 품질이 높지만 경계 hard constraint를 제공하지 않고 더 많은 AgentOutput을 전송했다."
)
heading(doc, "4.1 통계적 해석", 2)
para(
    doc,
    "40개 의미군 bootstrap에서 Calibrated Boundary와 Static의 Macro-F1 차이는 -2.91%p(95% CI -6.30~0.62), "
    "Top-k 대비 -1.98%p(-5.44~1.58)였다. 사전 설정한 5%p 비열등성 한계보다 CI 하한이 각각 1.30%p, 0.44%p 낮아 "
    "비열등성을 확증하지 못했다. 연구 기반 adapter 대비 평균 차이는 양수였으나 Holm 보정 후 우월성 역시 확증되지 않았다. "
    "따라서 본 결과는 실용적 trade-off와 Pareto 개선의 증거이지 보편적 정확도 우월성의 증거가 아니다."
)

page_break(doc)
heading(doc, "5. 논의")
heading(doc, "5.1 연구적 의미", 2)
para(
    doc,
    "첫째, 데이터 보호를 단순 byte 감소나 privacy score로 환원하지 않고 routing 가능 영역을 제한하는 실행 제약으로 두었다. "
    "둘째, 학습형 정확도 점수와 정책 판단을 분리함으로써 모델이 높은 확률을 출력해도 금지된 경계를 우회하지 못한다. 셋째, "
    "단일 가중합 최고점 대신 품질·통신량 Pareto 전선을 보고해 운영기관이 위험 허용도에 따라 지점을 선택할 수 있게 했다. "
    "이 차별점은 강·약 LLM 선택 중심의 RouteLLM·IRT-Router, collaboration 최적화 중심의 MasRouter, query privacy 중심의 "
    "RemoteRAG와 구별된다."
)
para(
    doc,
    "기존 연구가 주로 ‘어떤 모델을 호출할 것인가’ 또는 ‘여러 Agent를 어떻게 협업시킬 것인가’를 묻는다면, 본 연구는 ‘어떤 데이터가 "
    "어느 경계를 넘을 수 있는가’를 먼저 결정한 뒤 그 허용 영역 안에서 전문가를 선택한다. 사후 마스킹이나 출력 검사는 이미 발생한 원문 "
    "접근과 전송을 되돌릴 수 없지만, 사전 경계 제약은 금지된 실행 경로 자체를 생성하지 않는다. 따라서 본 연구의 노블티는 새로운 분류기 "
    "하나가 아니라 데이터 거버넌스를 라우팅 외부의 준수 항목이 아닌 실행 제약으로 결합한 데 있다."
)
para(
    doc,
    "또한 품질을 단독으로 보고하지 않고 fan-out 및 경계 전송량과 함께 제시해 ‘조금 더 높은 F1을 위해 얼마나 많은 Agent와 데이터를 "
    "노출해야 하는가’를 비교 가능하게 했다. 최고 Static 대비 2.9%p의 F1 차이를 감수하고 15.0%의 호출과 23.2%의 전송량을 줄인 "
    "결과는 절대적 우월성 주장이 아니라 위험 허용도에 따른 운영점 선택의 근거이다. 이는 정확도만으로는 드러나지 않는 공공·기업 AX의 "
    "배치 가치를 정량화한다."
)
heading(doc, "5.2 공공·기업 AX 적용", 2)
para(
    doc,
    "공공기관은 개인정보·기밀 질의를 보안·법무 Agent에 강제하고, 일반 공개정보는 기술 Agent만 호출하도록 정책을 구성할 수 "
    "있다. 기업은 부서별 vector store와 로컬 LLM을 유지하면서 Core에는 구조화 판단과 evidence handle만 보낼 수 있다. "
    "정책 버전, 선택 Agent, 거부 이유, 전송 byte를 기록하면 사후 감사와 책임 추적도 가능하다. Edge-Core가 동일 PC에 있는 "
    "프로토타입에서도 논리적 경계를 검증할 수 있고, 향후 실제 노드 분리 시 RTT·queueing·straggler 평가로 확장할 수 있다."
)
heading(doc, "5.3 한계", 2)
para(
    doc,
    "첫째, 240문항은 40개 파일럿 의미군에서 확장돼 독립 통계 단위가 40개뿐이다. 둘째, 상용 LLM 라벨러의 exact agreement가 "
    "18.8%로 낮고 85%가 조정을 거쳤다. 이는 역할 정답의 본질적 모호성과 동일 공급자 계열 편향을 시사한다. 셋째, 연구 기반 "
    "baseline은 AX adapter이며 원 논문 checkpoint 재현이 아니다. 넷째, byte는 구조화 payload proxy이고 실제 분산망의 "
    "암호화·프로토콜 overhead와 GPU-second를 포함하지 않는다. 다섯째, 5%p 비열등성은 확증되지 않았다."
)
heading(doc, "6. 결론")
para(
    doc,
    "본 논문은 공공·기업 AX의 복수 전문 Agent 환경에서 정확도 calibration과 데이터 경계 hard constraint를 결합한 "
    "전문가 라우터를 제안하였다. 226문항 평가에서 최고 Static baseline의 96.0% Micro-F1을 유지하면서 fan-out 15.0%, "
    "경계 전송량 23.2%를 줄였고 Pareto 전선에 위치하였다. 이는 ‘가장 정확한 router’가 아니라 ‘경계 정책을 보장하면서 "
    "경쟁 가능한 품질과 낮은 통신량을 제공하는 router’라는 실용적 기여를 뒷받침한다. 향후 연구에서는 신규 의미군 200개 "
    "이상, 전문 라벨러 독립 검증, 실제 Edge 노드 배치, 비열등성 확증 및 정책 제약 ablation을 수행할 예정이다."
)

heading(doc, "참고문헌")
references = [
    "[1] P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” NeurIPS, 2020.",
    "[2] J. Wang et al., “Mixture-of-Agents Enhances Large Language Model Capabilities,” arXiv:2406.04692, 2024.",
    "[3] Y. Yue et al., “MasRouter: Learning to Route LLMs for Multi-Agent Systems,” ACL, 2025.",
    "[4] I. Ong et al., “RouteLLM: Learning to Route LLMs with Preference Data,” arXiv:2406.18665, 2024.",
    "[5] W. Song et al., “IRT-Router: Effective and Interpretable Multi-LLM Routing via Item Response Theory,” 2025.",
    "[6] Q. J. Hu et al., “RouterBench: A Benchmark for Multi-LLM Routing System,” arXiv:2403.12031, 2024.",
    "[7] Y. Cheng et al., “RemoteRAG: A Privacy-Preserving LLM Cloud RAG Service,” Findings of ACL, pp. 3820-3837, 2025.",
    "[8] NIST, “Artificial Intelligence Risk Management Framework (AI RMF 1.0),” NIST AI 100-1, 2023.",
    "[9] L. Zheng et al., “Judging LLM-as-a-Judge with MT-Bench and Chatbot Arena,” NeurIPS Datasets and Benchmarks, 2023.",
    "[10] 한국지능정보사회진흥원, “공공부문 AI 도입·활용 가이드,” 2026.",
]
for item in references:
    para(doc, item, size=7.2, after=1.2, first=0)

doc.core_properties.title = "공공·기업 AX를 위한 경계 제약형 보정 전문가 라우팅"
doc.core_properties.author = "Jihoon Lim"
doc.core_properties.subject = "국내 학술대회 투고 원고"
doc.core_properties.keywords = "Multi-Agent RAG, Expert Routing, Data Boundary, Public AX"
doc.save(OUT)
print(OUT)
